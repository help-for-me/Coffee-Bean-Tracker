import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..crud import EXPORT_COLUMNS


def build_xlsx(rows: list[dict], insights: dict) -> bytes:
    workbook = Workbook()

    raw_sheet = workbook.active
    raw_sheet.title = "Raw Data"
    raw_sheet.append(EXPORT_COLUMNS)
    for row in raw_sheet[1]:
        row.font = Font(bold=True)
    for row in rows:
        raw_sheet.append([row.get(col) for col in EXPORT_COLUMNS])
    _autosize_columns(raw_sheet)

    summary_sheet = workbook.create_sheet("Summary")
    _write_summary(summary_sheet, insights)
    _autosize_columns(summary_sheet)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _autosize_columns(sheet) -> None:
    for col_idx, column_cells in enumerate(sheet.columns, start=1):
        longest = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(longest + 2, 40)


def _write_summary(sheet, insights: dict) -> None:
    # All-time numbers only - a backup report should show the complete
    # picture, not whichever recent/all-time toggle the Insights page
    # happened to be on when this was generated.
    row_num = 1

    def write_table(title: str, headers: list[str], data_rows: list[tuple]) -> None:
        nonlocal row_num
        sheet.cell(row=row_num, column=1, value=title).font = Font(bold=True, size=13)
        row_num += 1
        for col_idx, header in enumerate(headers, start=1):
            sheet.cell(row=row_num, column=col_idx, value=header).font = Font(bold=True)
        row_num += 1
        for data_row in data_rows:
            for col_idx, value in enumerate(data_row, start=1):
                sheet.cell(row=row_num, column=col_idx, value=value)
            row_num += 1
        row_num += 1  # blank row between tables

    write_table(
        "Score by month",
        ["Month", "Avg score", "Count"],
        [(r["month"], r["avg_score"], r["count"]) for r in insights["monthly_trend"]],
    )
    write_table(
        "Favourite processes",
        ["Process", "Avg score", "Adjusted score", "Count"],
        [
            (r["process"], r["avg_score"], r["adjusted_score"], r["count"])
            for r in insights["by_process"]["all_time"]["items"]
        ],
    )
    write_table(
        "Favourite origin countries",
        ["Origin country", "Avg score", "Adjusted score", "Count"],
        [
            (r["origin_country"], r["avg_score"], r["adjusted_score"], r["count"])
            for r in insights["by_origin_country"]["all_time"]["items"]
        ],
    )
    write_table(
        "Favourite tasting notes",
        ["Note", "Avg score", "Adjusted score", "Count"],
        [
            (r["note"], r["avg_score"], r["adjusted_score"], r["count"])
            for r in insights["by_tasting_note"]["all_time"]["items"]
        ],
    )
    write_table(
        "Most repurchased",
        ["Roaster", "Bean name", "Entry count", "Avg score", "Trend"],
        [
            (r["roaster"], r["bean_name"], r["entry_count"], r["avg_score"], r["trend"])
            for r in insights["most_repurchased"]["all_time"]
        ],
    )
