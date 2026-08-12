import io

from openpyxl import load_workbook

from backend.crud import EXPORT_COLUMNS
from backend.exports.csv_sink import build_csv
from backend.exports.xlsx_sink import build_xlsx


def _sample_row(**overrides):
    row = dict.fromkeys(EXPORT_COLUMNS)
    row.update(
        entry_id=1,
        rating_id=1,
        roaster="Stumptown",
        bean_name="Hair Bender",
        is_provisional=False,
        entry_type="bag",
        score=8.5,
        process="Washed",
    )
    row.update(overrides)
    return row


_NO_SIGNIFICANCE = {"comparable": False, "p_value": None, "significant": None, "message": "No data yet."}
_EMPTY_RANKING = {"items": [], "significance": _NO_SIGNIFICANCE}


def _sample_insights():
    return {
        "monthly_trend": [{"month": "2026-08", "avg_score": 8.0, "count": 1}],
        "by_process": {
            "all_time": {
                "items": [{"process": "Washed", "avg_score": 8.0, "adjusted_score": 8.0, "count": 1}],
                "significance": _NO_SIGNIFICANCE,
            },
            "recent": _EMPTY_RANKING,
        },
        "by_origin_country": {"all_time": _EMPTY_RANKING, "recent": _EMPTY_RANKING},
        "by_tasting_note": {"all_time": _EMPTY_RANKING, "recent": _EMPTY_RANKING},
        "most_repurchased": {"all_time": [], "recent": []},
    }


# --- CSV ---


def test_build_csv_header_matches_export_columns():
    csv_text = build_csv([])
    header_line = csv_text.strip().splitlines()[0]
    assert header_line.split(",") == EXPORT_COLUMNS


def test_build_csv_includes_row_data():
    csv_text = build_csv([_sample_row()])
    lines = csv_text.strip().splitlines()
    assert len(lines) == 2
    assert "Stumptown" in lines[1]
    assert "8.5" in lines[1]


def test_build_csv_empty_rows_still_has_header_only():
    csv_text = build_csv([])
    assert len(csv_text.strip().splitlines()) == 1


# --- XLSX ---


def test_build_xlsx_has_raw_data_and_summary_sheets():
    content = build_xlsx([_sample_row()], _sample_insights())
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Raw Data", "Summary"]


def test_build_xlsx_raw_data_header_matches_export_columns():
    content = build_xlsx([], _sample_insights())
    workbook = load_workbook(io.BytesIO(content))
    header = [cell.value for cell in workbook["Raw Data"][1]]
    assert header == EXPORT_COLUMNS


def test_build_xlsx_raw_data_includes_row_values():
    content = build_xlsx([_sample_row()], _sample_insights())
    sheet = load_workbook(io.BytesIO(content))["Raw Data"]
    roaster_col = EXPORT_COLUMNS.index("roaster")
    assert sheet[2][roaster_col].value == "Stumptown"


def test_build_xlsx_summary_includes_process_table():
    content = build_xlsx([], _sample_insights())
    sheet = load_workbook(io.BytesIO(content))["Summary"]
    values = {cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}
    assert "Favourite processes" in values
    assert "Washed" in values


def test_build_xlsx_empty_database_still_produces_valid_workbook():
    empty_insights = {
        "monthly_trend": [],
        "by_process": {"all_time": _EMPTY_RANKING, "recent": _EMPTY_RANKING},
        "by_origin_country": {"all_time": _EMPTY_RANKING, "recent": _EMPTY_RANKING},
        "by_tasting_note": {"all_time": _EMPTY_RANKING, "recent": _EMPTY_RANKING},
        "most_repurchased": {"all_time": [], "recent": []},
    }
    content = build_xlsx([], empty_insights)
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Raw Data", "Summary"]
